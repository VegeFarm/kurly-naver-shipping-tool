from __future__ import annotations

import hashlib
import io
import re
import zipfile
from dataclasses import dataclass
from typing import Iterable

from openpyxl import load_workbook


REQUIRED_HEADERS = {
    "통합배송지",
    "우편번호",
    "상품주문번호",
    "주문번호",
    "구매자연락처",
}


@dataclass
class OrderRow:
    row_number: int
    address: str
    zip_code: str
    product_order_id: str
    order_id: str
    buyer_phone: str

    @property
    def address_key(self) -> str:
        return f"{self.zip_code}|{normalize_address(self.address)}"

    @property
    def address_hash(self) -> str:
        return hashlib.sha256(self.address_key.encode("utf-8")).hexdigest()


def _cell_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize_address(value: str) -> str:
    return re.sub(r"\s+", " ", (value or "").strip())


def decrypt_excel_bytes(data: bytes, password: str) -> bytes:
    if zipfile.is_zipfile(io.BytesIO(data)):
        return data

    try:
        import msoffcrypto
    except ImportError as exc:  # pragma: no cover - deployment dependency
        raise RuntimeError("msoffcrypto-tool이 설치되어 있지 않습니다.") from exc

    src = io.BytesIO(data)
    out = io.BytesIO()
    try:
        office = msoffcrypto.OfficeFile(src)
        office.load_key(password=password, verify_password=True)
        office.decrypt(out, verify_integrity=True)
    except Exception as exc:
        raise ValueError("엑셀 비밀번호가 맞지 않거나 지원되지 않는 암호화 형식입니다.") from exc
    return out.getvalue()


def _find_order_sheet_and_header(wb):
    for ws in wb.worksheets:
        for row_idx in range(1, min(ws.max_row, 12) + 1):
            headers = {
                _cell_text(ws.cell(row_idx, col).value): col
                for col in range(1, ws.max_column + 1)
            }
            if REQUIRED_HEADERS.issubset(headers.keys()):
                return ws, row_idx, headers
    raise ValueError(
        "네이버 주문 엑셀의 필수 열(통합배송지/우편번호/상품주문번호/주문번호/구매자연락처)을 찾지 못했습니다."
    )


def parse_orders(plain_xlsx: bytes) -> tuple[list[OrderRow], dict]:
    wb = load_workbook(io.BytesIO(plain_xlsx), data_only=False, read_only=False, keep_links=True)
    ws, header_row, headers = _find_order_sheet_and_header(wb)

    rows: list[OrderRow] = []
    for r in range(header_row + 1, ws.max_row + 1):
        product_order_id = _cell_text(ws.cell(r, headers["상품주문번호"]).value)
        order_id = _cell_text(ws.cell(r, headers["주문번호"]).value)
        address = _cell_text(ws.cell(r, headers["통합배송지"]).value)
        zip_code = _cell_text(ws.cell(r, headers["우편번호"]).value)
        buyer_phone = _cell_text(ws.cell(r, headers["구매자연락처"]).value)
        if not product_order_id and not order_id and not address:
            continue
        if not product_order_id:
            continue
        rows.append(
            OrderRow(
                row_number=r,
                address=address,
                zip_code=zip_code,
                product_order_id=product_order_id,
                order_id=order_id or product_order_id,
                buyer_phone=buyer_phone,
            )
        )

    meta = {
        "sheet_name": ws.title,
        "header_row": header_row,
        "max_column": ws.max_column,
        "row_count": len(rows),
    }
    return rows, meta


def _encrypt_ooxml(plain_bytes: bytes, password: str) -> bytes:
    if not password:
        return plain_bytes
    try:
        from msoffcrypto.format.ooxml import OOXMLFile
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("msoffcrypto-tool이 설치되어 있지 않습니다.") from exc

    out = io.BytesIO()
    OOXMLFile(io.BytesIO(plain_bytes)).encrypt(password, out)
    return out.getvalue()


def build_filtered_workbook(
    plain_xlsx: bytes,
    keep_row_numbers: Iterable[int],
    output_password: str = "",
) -> bytes:
    """원본 workbook을 복제한 뒤 데이터 행만 제거하여 서식/폭/병합/필터를 최대한 그대로 유지한다."""
    keep = set(keep_row_numbers)
    wb = load_workbook(io.BytesIO(plain_xlsx), data_only=False, read_only=False, keep_links=True)
    ws, header_row, _ = _find_order_sheet_and_header(wb)

    # 아래에서 위로 삭제해야 행 번호 이동으로 인한 오삭제를 막을 수 있다.
    for r in range(ws.max_row, header_row, -1):
        if r not in keep:
            ws.delete_rows(r, 1)

    plain_out = io.BytesIO()
    wb.save(plain_out)
    return _encrypt_ooxml(plain_out.getvalue(), output_password)


def digits_only(phone: str) -> str:
    return re.sub(r"\D", "", phone or "")


def format_korean_phone(phone: str) -> str:
    d = digits_only(phone)
    if not d:
        return ""
    if d.startswith("02"):
        if len(d) == 9:
            return f"{d[:2]}-{d[2:5]}-{d[5:]}"
        if len(d) == 10:
            return f"{d[:2]}-{d[2:6]}-{d[6:]}"
    if len(d) == 11:
        return f"{d[:3]}-{d[3:7]}-{d[7:]}"
    if len(d) == 10:
        return f"{d[:3]}-{d[3:6]}-{d[6:]}"
    return d
